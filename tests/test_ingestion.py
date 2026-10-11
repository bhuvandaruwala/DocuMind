"""Loaders, chunking, vector store."""

import pytest

from app.ingest.chunking import chunk_documents, recursive_split
from app.ingest.loaders import Document, IngestionError, load_document


def test_pdf_has_pages(sample_dir):
    docs = load_document(sample_dir / "Nimbus_Employee_Handbook_2026.pdf")
    assert len(docs) == 3
    assert [d.metadata["page"] for d in docs] == [1, 2, 3]
    assert "24 days of paid annual leave" in docs[0].text


def test_txt_loads(sample_dir):
    docs = load_document(sample_dir / "nimbus_it_helpdesk_faq.txt")
    assert len(docs) == 1 and "GlobalProtect" in docs[0].text


def test_docx_text_and_tables(tmp_path):
    from docx import Document as WordDocument

    word = WordDocument()
    word.add_heading("IT Security Policy", level=1)
    word.add_paragraph("Passwords must be rotated every 90 days.")
    table = word.add_table(rows=2, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "Device", "Owner"
    table.cell(1, 0).text, table.cell(1, 1).text = "MacBook Pro", "Asha"
    word.add_paragraph("Report stolen laptops within 24 hours.")
    path = tmp_path / "policy.docx"
    word.save(path)

    docs = load_document(path)
    assert len(docs) == 1 and docs[0].metadata["file_type"] == "docx"
    text = docs[0].text
    assert "rotated every 90 days" in text and "MacBook Pro | Asha" in text
    # the table stays between the paragraphs around it
    assert text.index("90 days") < text.index("MacBook Pro") < text.index("24 hours")


def test_csv_rows_and_summary(sample_dir):
    docs = load_document(sample_dir / "nimbus_sales_q3_2026.csv")
    rows = [d for d in docs if d.metadata["kind"] == "table_rows"]
    summaries = [d for d in docs if d.metadata["kind"] == "table_summary"]
    assert rows and summaries
    assert all(d.text.startswith("Table: nimbus_sales_q3_2026.csv") for d in rows)  # header repeated
    joined = " ".join(d.text for d in summaries)
    assert "category = Electronics" in joined and "17,773" in joined
    assert "Highest total revenue_inr_lakh by region: West" in joined


def test_excel_all_sheets(sample_dir):
    docs = load_document(sample_dir / "nimbus_it_asset_inventory.xlsx")
    sheets = {d.metadata.get("sheet") for d in docs}
    assert sheets == {"Laptops", "Software Licenses"}
    assert any("In Repair: 2" in d.text for d in docs)


def test_bad_inputs(tmp_path):
    fake_docx = tmp_path / "broken.docx"
    fake_docx.write_bytes(b"not a word file")
    with pytest.raises(IngestionError):
        load_document(fake_docx)
    fake_pdf = tmp_path / "broken.pdf"
    fake_pdf.write_bytes(b"%PDF-1.4 not really a pdf")
    with pytest.raises(IngestionError):
        load_document(fake_pdf)
    empty = tmp_path / "empty.txt"
    empty.write_text("   ")
    with pytest.raises(IngestionError):
        load_document(empty)
    with pytest.raises(IngestionError):
        load_document(tmp_path / "x.docx")


def test_recursive_split_respects_size():
    text = ("Sentence number %d is here. " * 200) % tuple(range(200))
    chunks = recursive_split(text, 300, 50)
    assert len(chunks) > 5
    assert all(len(c) <= 300 for c in chunks)


def test_chunk_ids_unique_and_located():
    docs = [Document("para one.\n\n" + "word " * 400, {"source": "a.pdf", "page": 2, "file_type": "pdf"})]
    chunks = chunk_documents(docs, 400, 60)
    ids = [c["id"] for c in chunks]
    assert len(ids) == len(set(ids))
    assert chunks[0]["metadata"]["location"] == "a.pdf, page 2"


def test_store_persists_and_searches(loaded_store, tmp_path):
    from app.store import VectorStore

    assert loaded_store.count() > 30
    hits = loaded_store.keyword_search("GlobalProtect VPN portal", 3)
    assert hits[0]["metadata"]["source"] == "nimbus_it_helpdesk_faq.txt"
    # re-open from disk: data survives a restart
    reopened = VectorStore("hash", persist_dir=tmp_path / "chroma")
    assert reopened.count() == loaded_store.count()


def test_reingest_replaces_file(loaded_store, settings, sample_dir):
    from app.ingest.pipeline import ingest_folder

    before = loaded_store.count()
    ingest_folder(sample_dir, loaded_store, settings)
    assert loaded_store.count() == before


def test_delete_source(loaded_store):
    loaded_store.delete_source("vendor_meeting_notes.txt")
    assert "vendor_meeting_notes.txt" not in loaded_store.list_sources()
