"""
Document ingestion: PDF, Word, TXT, CSV and Excel loaders.

Every loader returns a list of `Document(text, metadata)`. Metadata always holds
`source` (file name) and `file_type`, plus location info used for citations:
  PDF   -> page
  Word  -> (none; chunk offsets are added later)
  TXT   -> (none; chunk offsets are added later)
  CSV   -> rows ("2-9", spreadsheet-style row numbers incl. header)
  Excel -> sheet + rows

For tables we also create *summary documents* (column statistics, totals per
category, largest/smallest values). Vector search over raw rows cannot answer
"what is the total revenue for the West region?", but a pre-computed summary can.
"""

from __future__ import annotations

import io
import logging
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Dict, List

import pandas as pd

log = logging.getLogger(__name__)


class IngestionError(Exception):
    """A single file could not be read (the batch continues with other files)."""


@dataclass
class Document:
    text: str
    metadata: Dict[str, Any] = field(default_factory=dict)


# PDF
def load_pdf(path: Path) -> List[Document]:
    from pypdf import PdfReader
    from pypdf.errors import PdfReadError

    try:
        reader = PdfReader(str(path))
    except (PdfReadError, OSError, ValueError) as e:
        raise IngestionError(f"{path.name}: not a readable PDF ({e})") from e
    if reader.is_encrypted:
        try:
            if reader.decrypt("") == 0:
                raise IngestionError(f"{path.name}: PDF is password-protected")
        except Exception as e:  # noqa: BLE001
            raise IngestionError(f"{path.name}: PDF is password-protected") from e

    docs: List[Document] = []
    for i, page in enumerate(reader.pages, start=1):
        try:
            text = page.extract_text() or ""
        except Exception as e:  # noqa: BLE001 - one bad page must not kill the file
            log.warning("%s page %s unreadable: %s", path.name, i, e)
            continue
        text = _clean(text)
        if text:
            docs.append(Document(text, {"source": path.name, "file_type": "pdf", "page": i}))
    if not docs:
        raise IngestionError(
            f"{path.name}: no extractable text (it may be a scanned image PDF; OCR is not supported)"
        )
    return docs


# Word
def load_docx(path: Path) -> List[Document]:
    from docx import Document as WordDocument
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    try:
        doc = WordDocument(str(path))
    except Exception as e:  # noqa: BLE001 - python-docx raises several types for bad files
        raise IngestionError(f"{path.name}: not a readable Word document ({e})") from e

    # walk the body in order so tables stay next to the text around them
    lines: List[str] = []
    for child in doc.element.body.iterchildren():
        if child.tag.endswith("}p"):
            text = Paragraph(child, doc).text.strip()
            if text:
                lines.append(text)
        elif child.tag.endswith("}tbl"):
            for row in Table(child, doc).rows:
                cells = []
                for cell in row.cells:
                    value = " ".join(cell.text.split())
                    if value and (not cells or cells[-1] != value):  # merged cells repeat their text
                        cells.append(value)
                if cells:
                    lines.append(" | ".join(cells))
    text = _clean("\n\n".join(lines))
    if not text:
        raise IngestionError(f"{path.name}: no text found in the document")
    return [Document(text, {"source": path.name, "file_type": "docx"})]


# TXT
def _read_text(path: Path) -> str:
    raw = path.read_bytes()
    for enc in ("utf-8-sig", "utf-8", "cp1252", "latin-1"):
        try:
            return raw.decode(enc)
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", errors="replace")


def load_text(path: Path) -> List[Document]:
    text = _clean(_read_text(path))
    if not text:
        raise IngestionError(f"{path.name}: file is empty")
    return [Document(text, {"source": path.name, "file_type": "txt"})]


# Tables
def _fmt(v: Any) -> str:
    if isinstance(v, float):
        if pd.isna(v):
            return ""
        return f"{v:,.2f}".rstrip("0").rstrip(".") if not v.is_integer() else f"{int(v):,}"
    if isinstance(v, (int,)) and not isinstance(v, bool):
        return f"{v:,}"
    if isinstance(v, pd.Timestamp):
        return v.date().isoformat()
    return str(v).strip()


def _clean_frame(df: pd.DataFrame) -> pd.DataFrame:
    df = df.dropna(how="all").dropna(axis=1, how="all")
    df.columns = [str(c).strip() or f"column_{i+1}" for i, c in enumerate(df.columns)]
    return df.reset_index(drop=True)


def table_to_documents(
    df: pd.DataFrame, source: str, file_type: str, sheet: str | None = None, rows_per_chunk: int = 8
) -> List[Document]:
    """Row-group documents (header repeated in every group) + summary documents."""
    df = _clean_frame(df)
    if df.empty:
        return []
    label = f"{source}" + (f" / sheet '{sheet}'" if sheet else "")
    base_meta = {"source": source, "file_type": file_type}
    if sheet:
        base_meta["sheet"] = sheet

    docs: List[Document] = []
    cols = list(df.columns)
    for start in range(0, len(df), rows_per_chunk):
        part = df.iloc[start : start + rows_per_chunk]
        first, last = start + 2, start + len(part) + 1  # +2: 1-based and header row
        lines = [f"Table: {label}. Columns: {', '.join(cols)}. Rows {first}-{last}:"]
        for _, row in part.iterrows():
            lines.append("; ".join(f"{c}: {_fmt(row[c])}" for c in cols if _fmt(row[c]) != ""))
        docs.append(Document("\n".join(lines), {**base_meta, "rows": f"{first}-{last}", "kind": "table_rows"}))

    docs.extend(summarize_table(df, label, base_meta))
    return docs


def summarize_table(df: pd.DataFrame, label: str, base_meta: Dict[str, Any]) -> List[Document]:
    """Pre-computed statistics so aggregate questions can be answered from retrieved text."""
    out: List[Document] = []
    num_cols = [c for c in df.columns if pd.api.types.is_numeric_dtype(df[c]) and not pd.api.types.is_bool_dtype(df[c])]
    cat_cols = [
        c for c in df.columns
        if c not in num_cols and 1 < df[c].nunique(dropna=True) <= 25 and df[c].nunique() < len(df)
    ]
    id_col = next((c for c in df.columns if c not in num_cols and df[c].nunique() == len(df)), None)

    # 1) overview
    lines = [f"Summary of table {label}: {len(df)} data rows and {len(df.columns)} columns ({', '.join(df.columns)})."]
    for c in num_cols:
        s = df[c].dropna()
        if s.empty:
            continue
        lines.append(
            f"Column {c}: total (sum) {_fmt(float(s.sum()))}, average {_fmt(round(float(s.mean()), 2))}, "
            f"minimum {_fmt(float(s.min()))}, maximum {_fmt(float(s.max()))}."
        )
        if id_col is not None:
            hi, lo = df.loc[s.idxmax()], df.loc[s.idxmin()]
            lines.append(
                f"Highest {c}: {hi[id_col]} ({_fmt(float(hi[c]))}). Lowest {c}: {lo[id_col]} ({_fmt(float(lo[c]))})."
            )
    for c in cat_cols:
        vc = df[c].value_counts()
        lines.append(f"Count of rows by {c}: " + ", ".join(f"{k}: {v}" for k, v in vc.items()) + ".")
    out.append(Document("\n".join(lines), {**base_meta, "rows": "all", "kind": "table_summary"}))

    # 2) totals per category column (one small document per categorical column)
    for c in cat_cols:
        if not num_cols:
            break
        g = df.groupby(c, dropna=True)[num_cols].sum(numeric_only=True)
        lines = [f"Totals by {c} for table {label} (sum of each numeric column per {c}):"]
        for key, row in g.iterrows():
            lines.append(f"{c} = {key}: " + "; ".join(f"total {n} {_fmt(float(row[n]))}" for n in num_cols) + ".")
        for n in num_cols:
            top = g[n].sort_values(ascending=False)
            lines.append(f"Highest total {n} by {c}: {top.index[0]} ({_fmt(float(top.iloc[0]))}); "
                         f"lowest: {top.index[-1]} ({_fmt(float(top.iloc[-1]))}).")
        out.append(Document("\n".join(lines), {**base_meta, "rows": "all", "kind": "table_summary"}))
    return out


def _read_csv(path: Path) -> pd.DataFrame:
    text = _read_text(path)
    try:
        return pd.read_csv(io.StringIO(text), sep=None, engine="python")
    except Exception:  # noqa: BLE001
        return pd.read_csv(io.StringIO(text))


def load_csv(path: Path, rows_per_chunk: int = 8) -> List[Document]:
    try:
        df = _read_csv(path)
    except Exception as e:  # noqa: BLE001
        raise IngestionError(f"{path.name}: could not parse CSV ({e})") from e
    docs = table_to_documents(df, path.name, "csv", rows_per_chunk=rows_per_chunk)
    if not docs:
        raise IngestionError(f"{path.name}: CSV has no data rows")
    return docs


def load_excel(path: Path, rows_per_chunk: int = 8) -> List[Document]:
    try:
        sheets = pd.read_excel(path, sheet_name=None)
    except Exception as e:  # noqa: BLE001
        raise IngestionError(f"{path.name}: could not open Excel file ({e})") from e
    docs: List[Document] = []
    for name, df in sheets.items():
        docs.extend(table_to_documents(df, path.name, "excel", sheet=str(name), rows_per_chunk=rows_per_chunk))
    if not docs:
        raise IngestionError(f"{path.name}: all sheets are empty")
    return docs


# Router
def load_document(path: str | Path, rows_per_chunk: int = 8) -> List[Document]:
    path = Path(path)
    ext = path.suffix.lower()
    if ext == ".pdf":
        return load_pdf(path)
    if ext == ".docx":
        return load_docx(path)
    if ext == ".txt":
        return load_text(path)
    if ext == ".csv":
        return load_csv(path, rows_per_chunk)
    if ext in (".xlsx", ".xls"):
        return load_excel(path, rows_per_chunk)
    raise IngestionError(f"{path.name}: unsupported file type '{ext}'")


def _clean(text: str) -> str:
    text = text.replace("\x00", "").replace("\r\n", "\n").replace("\r", "\n")
    lines = [" ".join(line.split()) for line in text.split("\n")]
    out, blank = [], 0
    for line in lines:
        if not line:
            blank += 1
            if blank <= 1:
                out.append("")
        else:
            blank = 0
            out.append(line)
    return "\n".join(out).strip()
